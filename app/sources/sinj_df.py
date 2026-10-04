"""Retrieve and identity-check legislative text from the official SINJ-DF."""
from __future__ import annotations

import json
import html
import io
import re
import urllib.parse
import urllib.request
from dataclasses import dataclass
from datetime import date, datetime
from urllib.error import HTTPError

from bs4 import BeautifulSoup

from app.catalog_sync.sinj_df import SINJ_WEB
from app.sources.history import OfficialRelation
from app.sources.normas import SourceDocumentUnavailable


MAX_DOCUMENT_BYTES = 25_000_000
REMOTE_ID = re.compile(r"(?:[0-9a-f]{32}|[0-9]{1,24})")


@dataclass(frozen=True)
class SinjDFDocument:
    body: bytes
    source_url: str
    representation: str
    version: str
    parsed_body: bytes | None = None
    raw_content_type: str = "text/html"
    legal_value: str = "UnclassifiedLegalValue"
    notice: str = "Texto recuperado do portal oficial SINJ-DF; sua classificação jurídica não é informada pelo portal."


@dataclass(frozen=True)
class SinjDFHistorySnapshot:
    body: bytes
    source_url: str
    relations: list[OfficialRelation]


def _fetch(url: str, *, timeout: int) -> tuple[bytes, str, str]:
    request = urllib.request.Request(url, headers={
        "Accept": "text/html, application/xhtml+xml",
        "Referer": f"{SINJ_WEB}/ResultadoDePesquisa?tipo_pesquisa=norma",
        "User-Agent": "LeiAberta/0.3 (+fontes oficiais)",
    })
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            body = response.read(MAX_DOCUMENT_BYTES + 1)
            if response.status != 200 or len(body) > MAX_DOCUMENT_BYTES:
                raise SourceDocumentUnavailable("A resposta do SINJ-DF falhou ou excedeu o limite de tamanho permitido.")
            return body, response.geturl(), response.headers.get("Content-Type", "")
    except HTTPError as exc:
        if exc.code in {400, 404, 410, 422}:
            raise SourceDocumentUnavailable(f"O SINJ-DF não disponibiliza este documento (HTTP {exc.code}).") from exc
        raise


def _record_id(source_url: str) -> str:
    parsed = urllib.parse.urlparse(source_url)
    query = urllib.parse.parse_qs(parsed.query)
    document_ids = query.get("id_doc", [])
    if (parsed.scheme != "https" or parsed.hostname != "www.sinj.df.gov.br"
            or parsed.path != "/sinj/DetalhesDeNorma.aspx" or len(document_ids) != 1
            or not document_ids[0].isdigit()):
        raise SourceDocumentUnavailable("A URL não aponta para um registro individual reconhecido do SINJ-DF.")
    return document_ids[0]


def _detail_metadata(body: bytes) -> dict:
    soup = BeautifulSoup(body, "html.parser")
    script = next((tag.string or tag.get_text() for tag in soup.find_all("script")
                   if "var json_norma" in (tag.string or tag.get_text())), None)
    if not script:
        raise SourceDocumentUnavailable("A ficha SINJ-DF não inclui metadados estruturados da norma.")
    match = re.search(r"var\s+json_norma\s*=\s*(\{.*?\})\s*;", script, re.S)
    if not match:
        raise SourceDocumentUnavailable("Os metadados da ficha SINJ-DF não puderam ser interpretados.")
    try:
        metadata = json.loads(match.group(1))
    except json.JSONDecodeError as exc:
        raise SourceDocumentUnavailable("Os metadados da ficha SINJ-DF não são JSON válido.") from exc
    if not isinstance(metadata, dict):
        raise SourceDocumentUnavailable("A ficha SINJ-DF não retornou uma norma individual.")
    return metadata


def _expected_number(value: str) -> str:
    return re.sub(r"[^\d-]", "", value)


def _date(value: object) -> date | None:
    if not value:
        return None
    try:
        return datetime.strptime(str(value).strip(), "%d/%m/%Y").date()
    except ValueError:
        return None


def _filename(metadata: dict) -> tuple[str, str]:
    """Match the public filename produced by SINJ-DF's getTitleNorma function."""
    updated = metadata.get("ar_atualizado") if isinstance(metadata.get("ar_atualizado"), dict) else {}
    sources = metadata.get("fontes") if isinstance(metadata.get("fontes"), list) else []
    selected = updated if updated.get("id_file") else None
    version = "Current" if selected else "Original"
    if selected is None:
        publication_file = None
        reissue_file = None
        for item in sources:
            if not isinstance(item, dict) or not isinstance(item.get("ar_fonte"), dict):
                continue
            attached = item["ar_fonte"]
            if not attached.get("id_file"):
                continue
            publication_type = str(item.get("nm_tipo_publicacao") or "").strip().casefold()
            if publication_type in {"publicação", "pub"}:
                publication_file = attached
            if publication_type in {"republicação", "rep", "retificação", "ret"}:
                reissue_file = attached
                version = "Republished"
                break
        selected = reissue_file or publication_file
    if not selected:
        raise SourceDocumentUnavailable("A ficha SINJ-DF não possui arquivo integral anexado à norma.")
    mimetype = str(selected.get("mimetype") or "").casefold()
    if "html" in mimetype or "htm" in mimetype:
        extension = ".html"
    elif "pdf" in mimetype:
        extension = ".pdf"
    elif "wordprocessingml.document" in mimetype or mimetype.endswith("/msword"):
        extension = ".docx" if "wordprocessingml" in mimetype else ".doc"
    else:
        raise SourceDocumentUnavailable(f"O SINJ-DF oferece anexo em formato não suportado ({mimetype or 'desconhecido'}).")
    law_type = str(metadata.get("nm_tipo_norma") or "").strip()
    number = str(metadata.get("nr_norma") or "").strip()
    date = str(metadata.get("dt_assinatura") or "").strip()
    display_number = "" if number == "0" else number
    basename = f"{law_type} {display_number}_{date}"
    # JavaScript's \W is ASCII-based. The SINJ page applies the same rule to
    # punctuation and Portuguese diacritics before appending the extension.
    basename = re.sub(r"\W+|[ãÃõÕçÇêÊéÉ/]", "_", basename, flags=re.ASCII)
    return f"{basename}{extension}", version


def _validate_html(body: bytes, content_type: str, law_type: str, number: str, signed: datetime) -> bytes:
    if "text/html" not in content_type.casefold() or b"<html" not in body[:8192].lower():
        raise SourceDocumentUnavailable("O SINJ-DF não retornou HTML de texto integral nesta ficha.")
    soup = BeautifulSoup(body, "html.parser")
    actual = " ".join((soup.title.get_text(" ", strip=True) if soup.title else "").split())
    expected = f"{law_type} {number} de {signed:%d/%m/%Y}".strip()
    if number and number != "0" and expected.casefold() not in actual.casefold():
        raise SourceDocumentUnavailable("O texto baixado do SINJ-DF não identifica o tipo, número e data esperados.")
    visible = soup.get_text(" ", strip=True)
    if len(visible) < 80 or "norma não encontrada" in visible.casefold():
        raise SourceDocumentUnavailable("A URL do SINJ-DF não contém texto legislativo legível.")
    return body


def _as_paragraph_html(text: str) -> bytes:
    paragraphs = [line.strip() for line in text.splitlines() if line.strip()]
    if not paragraphs or sum(map(len, paragraphs)) < 80:
        raise SourceDocumentUnavailable("O anexo SINJ-DF não contém texto integral legível, mesmo após extração.")
    rendered = "\n".join(f"<p>{html.escape(line)}</p>" for line in paragraphs)
    return f"<!doctype html><html><body>{rendered}</body></html>".encode("utf-8")


def _extract_pdf(body: bytes) -> bytes:
    if not body.startswith(b"%PDF-"):
        raise SourceDocumentUnavailable("O SINJ-DF identificou PDF, mas os bytes recebidos não são um arquivo PDF.")
    try:
        import pymupdf

        document = pymupdf.open(stream=body, filetype="pdf")
        parts = []
        for page in document:
            page_text = page.get_text("text").strip()
            # A scanned page can contain a page number or a tiny header in its
            # text layer. OCR only pages that also contain images and little
            # selectable text, preserving the source PDF as the archived body.
            if len(page_text) < 80 and page.get_images(full=True):
                try:
                    ocr = page.get_textpage_ocr(language="por+eng", dpi=200, full=True)
                    ocr_text = page.get_text("text", textpage=ocr).strip()
                    if len(ocr_text) > len(page_text):
                        page_text = ocr_text
                except Exception as exc:
                    raise SourceDocumentUnavailable(
                        "O SINJ-DF publicou um PDF digitalizado, mas o OCR em português não pôde ser executado."
                    ) from exc
            if page_text:
                parts.append(page_text)
        document.close()
    except SourceDocumentUnavailable:
        raise
    except Exception as exc:
        raise SourceDocumentUnavailable(f"O PDF oficial do SINJ-DF não pôde ser extraído: {exc}") from exc
    return _as_paragraph_html("\n".join(parts))


def _extract_docx(body: bytes) -> bytes:
    try:
        from docx import Document

        document = Document(io.BytesIO(body))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    except Exception as exc:
        raise SourceDocumentUnavailable(f"O anexo DOCX oficial do SINJ-DF não pôde ser extraído: {exc}") from exc
    return _as_paragraph_html("\n".join(paragraphs))


def _extract_attachment(body: bytes, content_type: str, filename: str,
                        law_type: str, number: str, signed: datetime) -> tuple[bytes, str]:
    content_type = content_type.casefold()
    if filename.endswith(".html"):
        return _validate_html(body, content_type, law_type, number, signed), "text/html"
    if filename.endswith(".pdf") or "application/pdf" in content_type:
        return _extract_pdf(body), "application/pdf"
    if filename.endswith(".docx") or "wordprocessingml.document" in content_type:
        return _extract_docx(body), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    raise SourceDocumentUnavailable(f"O SINJ-DF oferece anexo em formato não suportado ({content_type or filename}).")


def fetch_sinj_df_document(
    source_url: str,
    law_type: str,
    number: str,
    year: int,
    *,
    timeout: int = 30,
) -> SinjDFDocument:
    """Resolve a page's own attached full-text URL and rejects identity drift."""
    document_id = _record_id(source_url)
    try:
        detail_body, detail_url, _detail_type = _fetch(source_url, timeout=timeout)
    except SourceDocumentUnavailable:
        raise
    except Exception as exc:
        raise SourceDocumentUnavailable(f"A ficha SINJ-DF não pôde ser lida: {exc}") from exc
    detail_location = urllib.parse.urlparse(detail_url)
    returned_document_id = urllib.parse.parse_qs(detail_location.query).get("id_doc", [])
    if (detail_location.scheme != "https" or detail_location.hostname != "www.sinj.df.gov.br"
            or detail_location.path != "/sinj/DetalhesDeNorma.aspx"
            or returned_document_id != [document_id]):
        raise SourceDocumentUnavailable("A ficha SINJ-DF redirecionou para domínio ou caminho não reconhecido.")
    metadata = _detail_metadata(detail_body)
    remote_hash = str(metadata.get("ch_norma") or "").strip().lower()
    signed_text = str(metadata.get("dt_assinatura") or "").strip()
    try:
        signed = datetime.strptime(signed_text, "%d/%m/%Y")
    except ValueError as exc:
        raise SourceDocumentUnavailable("A ficha SINJ-DF não informa uma data de assinatura válida.") from exc
    actual_type = str(metadata.get("nm_tipo_norma") or "").strip()
    actual_number = str(metadata.get("nr_norma") or "").strip()
    expected_number = "" if number.strip().casefold() in {"", "s/n", "0"} else _expected_number(number)
    actual_number_normalized = "" if actual_number in {"", "0"} else _expected_number(actual_number)
    if (not REMOTE_ID.fullmatch(remote_hash) or actual_type.casefold() != law_type.strip().casefold()
            or actual_number_normalized != expected_number or signed.year != year):
        raise SourceDocumentUnavailable(
            f"A identidade SINJ-DF diverge da norma pedida: {actual_type} {actual_number or 's/n'}/{signed.year}."
        )
    filename, version = _filename(metadata)
    text_url = f"{SINJ_WEB}/Norma/{remote_hash}/{urllib.parse.quote(filename, safe='._-')}"
    try:
        body, final_url, content_type = _fetch(text_url, timeout=timeout)
    except SourceDocumentUnavailable:
        raise
    except Exception as exc:
        raise SourceDocumentUnavailable(f"O texto SINJ-DF não pôde ser lido: {exc}") from exc
    location = urllib.parse.urlparse(final_url)
    expected_path = f"/sinj/Norma/{remote_hash}/{urllib.parse.quote(filename, safe='._-')}"
    if (location.scheme != "https" or location.hostname != "www.sinj.df.gov.br"
            or location.path != expected_path):
        raise SourceDocumentUnavailable("O texto SINJ-DF redirecionou para domínio ou caminho não reconhecido.")
    parsed_body, attachment_type = _extract_attachment(
        body, content_type, filename, law_type, actual_number, signed,
    )
    representation = "Texto integral disponibilizado pelo SINJ-DF"
    if version == "Current":
        representation = "Texto atualizado disponibilizado pelo SINJ-DF"
    elif version == "Republished":
        representation = "Texto republicado disponibilizado pelo SINJ-DF"
    if attachment_type == "application/pdf":
        representation += " (PDF oficial com texto extraído/OCR quando necessário)"
    elif attachment_type.startswith("application/vnd."):
        representation += " (DOCX oficial convertido para leitura)"
    return SinjDFDocument(body, final_url, representation, version, parsed_body, content_type)


def fetch_sinj_df_history(
    source_url: str,
    law_type: str,
    number: str,
    year: int,
    *,
    timeout: int = 30,
) -> SinjDFHistorySnapshot:
    """Read the portal's official incoming alteration and repeal relations."""
    document_id = _record_id(source_url)
    try:
        body, final_url, _content_type = _fetch(source_url, timeout=timeout)
    except SourceDocumentUnavailable:
        raise
    except Exception as exc:
        raise SourceDocumentUnavailable(f"A ficha SINJ-DF não pôde ser lida: {exc}") from exc
    location = urllib.parse.urlparse(final_url)
    returned_id = urllib.parse.parse_qs(location.query).get("id_doc", [])
    if (location.scheme != "https" or location.hostname != "www.sinj.df.gov.br"
            or location.path != "/sinj/DetalhesDeNorma.aspx" or returned_id != [document_id]):
        raise SourceDocumentUnavailable("A ficha SINJ-DF redirecionou para domínio, caminho ou norma diferente.")
    metadata = _detail_metadata(body)
    signed_text = str(metadata.get("dt_assinatura") or "").strip()
    try:
        signed = datetime.strptime(signed_text, "%d/%m/%Y")
    except ValueError as exc:
        raise SourceDocumentUnavailable("A ficha SINJ-DF não informa uma data de assinatura válida.") from exc
    actual_type = str(metadata.get("nm_tipo_norma") or "").strip()
    actual_number = str(metadata.get("nr_norma") or "").strip()
    expected_number = "" if number.strip().casefold() in {"", "s/n", "0"} else _expected_number(number)
    actual_number_normalized = "" if actual_number in {"", "0"} else _expected_number(actual_number)
    source_identifier = str(metadata.get("ch_norma") or "").strip().lower()
    if (not REMOTE_ID.fullmatch(source_identifier) or actual_type.casefold() != law_type.strip().casefold()
            or actual_number_normalized != expected_number or signed.year != year):
        raise SourceDocumentUnavailable(
            f"A identidade SINJ-DF diverge da norma pedida: {actual_type} {actual_number or 's/n'}/{signed.year}."
        )

    relations = []
    vides = metadata.get("vides")
    if not isinstance(vides, list):
        vides = []
    for index, item in enumerate(vides):
        if not isinstance(item, dict) or item.get("in_norma_afetada") is not True:
            continue
        action = str(item.get("ds_texto_relacao") or "").strip()
        relation_type = str(item.get("nm_tipo_relacao") or "").strip()
        normalized_action = f"{action} {relation_type}".casefold()
        if not any(word in normalized_action for word in (
            "altera", "acrescent", "revog", "restaur", "repristin", "retific", "nova redação", "derroga",
        )):
            continue
        relation_id = str(item.get("ch_vide") or f"{source_identifier}-{index}").strip().lower()
        amendment_id = str(item.get("ch_norma_vide") or "").strip().lower()
        if not REMOTE_ID.fullmatch(relation_id) or not REMOTE_ID.fullmatch(amendment_id):
            continue
        amendment_type = str(item.get("nm_tipo_norma_vide") or "Norma relacionada").strip()
        amendment_number = str(item.get("nr_norma_vide") or "").strip()
        amendment_date_text = str(item.get("dt_assinatura_norma_vide") or "").strip()
        amendment_date = None
        if amendment_date_text:
            try:
                amendment_date = datetime.strptime(amendment_date_text, "%d/%m/%Y").date()
            except ValueError:
                pass
        event_label = f"{amendment_type} {amendment_number}".strip()
        if amendment_date_text:
            event_label += f" de {amendment_date_text}"
        devices = []
        for key, label in (("artigo_norma_vide", "art."), ("paragrafo_norma_vide", "§"),
                           ("inciso_norma_vide", "inciso"), ("alinea_norma_vide", "alínea"),
                           ("anexo_norma_vide", "anexo")):
            value = item.get(key)
            if isinstance(value, str) and value.strip():
                devices.append(f"{label} {value.strip()}")
        change = item.get("alteracao_texto_vide")
        if isinstance(change, dict) and change.get("ds_dispositivos_alterados"):
            devices.append(str(change["ds_dispositivos_alterados"]).strip())
        evidence = f"SINJ-DF: {action or relation_type}; relação marcada como incidente sobre esta norma."
        if devices:
            evidence += " Dispositivos cadastrados: " + "; ".join(devices)
        comment = str(item.get("ds_comentario_vide") or "").strip()
        if comment:
            evidence += f" Observação do cadastro: {comment}"
        event_url = f"{SINJ_WEB}/DetalhesDeNorma.aspx?id_norma={urllib.parse.quote(amendment_id, safe='')}"
        relations.append(OfficialRelation(
            source_id=f"sinj:{relation_id}", device_ref="",
            relation=action or relation_type, event_label=event_label[:240], event_url=event_url,
            signed_at=amendment_date, publication_date=_date(item.get("dt_publicacao_fonte_norma_vide")),
            evidence=evidence[:4000],
        ))
    unique = {(item.source_id, item.event_url): item for item in relations}
    relations = sorted(unique.values(), key=lambda item: item.signed_at or date.min, reverse=True)
    return SinjDFHistorySnapshot(body=body, source_url=final_url, relations=relations)
