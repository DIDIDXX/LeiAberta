"""Extract text from official document attachments while retaining raw bytes."""
from __future__ import annotations

import html
import io

from app.sources.normas import SourceDocumentUnavailable


def text_to_html(text: str) -> bytes:
    paragraphs = [line.strip() for line in text.splitlines() if line.strip()]
    if not paragraphs or sum(map(len, paragraphs)) < 80:
        raise SourceDocumentUnavailable("O anexo oficial não contém texto integral legível após extração.")
    rendered = "\n".join(f"<p>{html.escape(line)}</p>" for line in paragraphs)
    return f"<!doctype html><html><body>{rendered}</body></html>".encode("utf-8")


def pdf_to_html(body: bytes) -> bytes:
    if not body.startswith(b"%PDF-"):
        raise SourceDocumentUnavailable("A fonte identificou PDF, mas os bytes recebidos não são um arquivo PDF.")
    try:
        import pymupdf

        document = pymupdf.open(stream=body, filetype="pdf")
        parts = []
        for page in document:
            page_text = page.get_text("text").strip()
            if len(page_text) < 80 and page.get_images(full=True):
                try:
                    ocr = page.get_textpage_ocr(language="por+eng", dpi=200, full=True)
                    ocr_text = page.get_text("text", textpage=ocr).strip()
                    if len(ocr_text) > len(page_text):
                        page_text = ocr_text
                except Exception as exc:
                    raise SourceDocumentUnavailable(
                        "O PDF oficial é digitalizado e o OCR em português não pôde ser executado."
                    ) from exc
            if page_text:
                parts.append(page_text)
        document.close()
    except SourceDocumentUnavailable:
        raise
    except Exception as exc:
        raise SourceDocumentUnavailable(f"O PDF oficial não pôde ser extraído: {exc}") from exc
    return text_to_html("\n".join(parts))


def docx_to_html(body: bytes) -> bytes:
    try:
        from docx import Document

        document = Document(io.BytesIO(body))
        paragraphs = [item.text.strip() for item in document.paragraphs if item.text.strip()]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    except Exception as exc:
        raise SourceDocumentUnavailable(f"O anexo DOCX oficial não pôde ser extraído: {exc}") from exc
    return text_to_html("\n".join(paragraphs))
