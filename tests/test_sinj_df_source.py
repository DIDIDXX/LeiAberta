import hashlib
import io
import json
import sys
import types
from pathlib import Path

import pytest

from app.sources.normas import SourceDocumentUnavailable
from app.sources.sinj_df import _extract_docx, _extract_pdf, fetch_sinj_df_document, fetch_sinj_df_history


FIXTURES = Path(__file__).parent / "fixtures" / "official"


def test_sinj_df_fetches_identity_checked_full_text(monkeypatch):
    manifest = json.loads((FIXTURES / "sinj-df-lei-6139-manifest.json").read_text())
    detail = (FIXTURES / "sinj-df-detail-lei-6139.html").read_bytes()
    body = (FIXTURES / "sinj-df-lei-6139.html").read_bytes()
    for filename, spec in manifest["files"].items():
        fixture = (FIXTURES / filename).read_bytes()
        assert len(fixture) == spec["bytes"]
        assert hashlib.sha256(fixture).hexdigest() == spec["sha256"]
    calls = []

    class Response(io.BytesIO):
        status = 200

        def __init__(self, content, url, content_type):
            super().__init__(content)
            self.url = url
            self.headers = {"Content-Type": content_type}

        def geturl(self):
            return self.url

    def fake_urlopen(request, timeout=25):
        url = request.full_url
        calls.append(url)
        if url == manifest["urls"]["detail"]:
            return Response(detail, url, "text/html; charset=utf-8")
        if url == manifest["urls"]["text"]:
            return Response(body, url, "text/html; charset=utf-8")
        raise AssertionError(f"URL inesperada: {url}")

    monkeypatch.setattr("app.sources.sinj_df.urllib.request.urlopen", fake_urlopen)
    document = fetch_sinj_df_document(
        manifest["urls"]["detail"], "Lei", "6139", 2018,
    )
    assert document.version == "Original"
    assert document.legal_value == "UnclassifiedLegalValue"
    assert document.source_url == manifest["urls"]["text"]
    assert "SINJ-DF" in document.representation
    assert len(calls) == 2


def test_sinj_df_rejects_a_detail_record_with_different_identity(monkeypatch):
    detail = (FIXTURES / "sinj-df-detail-lei-6139.html").read_bytes()
    detail = detail.replace(b'"nr_norma":"6139"', b'"nr_norma":"6140"', 1)

    class Response(io.BytesIO):
        status = 200
        headers = {"Content-Type": "text/html; charset=utf-8"}

        def geturl(self):
            return "https://www.sinj.df.gov.br/sinj/DetalhesDeNorma.aspx?id_doc=93048"

    monkeypatch.setattr("app.sources.sinj_df.urllib.request.urlopen", lambda *_a, **_kw: Response(detail))
    with pytest.raises(SourceDocumentUnavailable, match="identidade SINJ-DF diverge"):
        fetch_sinj_df_document(
            "https://www.sinj.df.gov.br/sinj/DetalhesDeNorma.aspx?id_doc=93048", "Lei", "6139", 2018,
        )


def test_sinj_df_resolves_legacy_numeric_norma_ids_and_updated_text(monkeypatch):
    manifest = json.loads((FIXTURES / "sinj-df-portaria-65-2015-manifest.json").read_text())
    detail = (FIXTURES / "sinj-df-detail-portaria-65-2015.html").read_bytes()
    body = (FIXTURES / "sinj-df-portaria-65-2015.html").read_bytes()
    calls = []

    class Response(io.BytesIO):
        status = 200

        def __init__(self, content, url):
            super().__init__(content)
            self.url = url
            self.headers = {"Content-Type": "text/html; charset=utf-8"}

        def geturl(self):
            return self.url

    def fake_urlopen(request, timeout=25):
        url = request.full_url
        calls.append(url)
        if url == manifest["urls"]["detail"]:
            return Response(detail, url)
        if url == manifest["urls"]["text"]:
            return Response(body, url)
        raise AssertionError(f"URL inesperada: {url}")

    monkeypatch.setattr("app.sources.sinj_df.urllib.request.urlopen", fake_urlopen)
    document = fetch_sinj_df_document(manifest["urls"]["detail"], "Portaria", "65", 2015)
    assert document.version == "Current"
    assert document.source_url == manifest["urls"]["text"]
    assert "atualizado" in document.representation
    assert len(calls) == 2


def test_sinj_df_history_reads_incoming_official_amendment_relations(monkeypatch):
    manifest = json.loads((FIXTURES / "sinj-df-decreto-36222-history-manifest.json").read_text())
    detail = (FIXTURES / "sinj-df-detail-decreto-36222-2014.html").read_bytes()
    spec = manifest["files"]["sinj-df-detail-decreto-36222-2014.html"]
    assert len(detail) == spec["bytes"]
    assert hashlib.sha256(detail).hexdigest() == spec["sha256"]

    class Response(io.BytesIO):
        status = 200
        headers = {"Content-Type": "text/html; charset=utf-8"}

        def geturl(self):
            return manifest["urls"]["detail"]

    monkeypatch.setattr("app.sources.sinj_df.urllib.request.urlopen", lambda *_a, **_kw: Response(detail))
    snapshot = fetch_sinj_df_history(manifest["urls"]["detail"], "Decreto", "36.222", 2014)
    assert len(snapshot.relations) >= 100
    assert snapshot.relations[0].source_id.startswith("sinj:")
    assert "id_norma=" in snapshot.relations[0].event_url
    assert "Alteração" in snapshot.relations[0].relation or "Alterado" in snapshot.relations[0].relation
    assert "relação marcada como incidente" in snapshot.relations[0].evidence


def test_sinj_df_extracts_searchable_text_from_official_pdf():
    import pymupdf

    source = pymupdf.open()
    page = source.new_page()
    page.insert_textbox(
        pymupdf.Rect(40, 40, 550, 800),
        "Art. 1º Esta Lei dispõe sobre a organização do serviço público no Distrito Federal.\n"
        "§ 1º A autoridade competente publicará as regras necessárias à execução desta Lei.\n"
        "Art. 2º Esta Lei entra em vigor na data de sua publicação.",
    )
    raw_pdf = source.tobytes()
    source.close()
    parsed = _extract_pdf(raw_pdf).decode("utf-8")
    assert "Art. 1º" in parsed
    assert "Art. 2º" in parsed


def test_sinj_df_uses_portuguese_ocr_for_scanned_pdf_pages(monkeypatch):
    expected = "Art. 1º Esta Lei estabelece regras administrativas no Distrito Federal. " * 3

    class Page:
        def get_text(self, kind, textpage=None):
            return expected if textpage is not None else "Página 1"

        def get_images(self, full=True):
            return [(1,)]

        def get_textpage_ocr(self, language, dpi, full):
            assert language == "por+eng"
            assert dpi == 200 and full is True
            return object()

    class Document(list):
        def close(self):
            pass

    monkeypatch.setitem(sys.modules, "pymupdf", types.SimpleNamespace(open=lambda **_: Document([Page()])))
    parsed = _extract_pdf(b"%PDF-scanned").decode("utf-8")
    assert "Art. 1º" in parsed
    assert "Distrito Federal" in parsed


def test_sinj_df_extracts_paragraphs_and_tables_from_official_docx():
    from docx import Document

    source = Document()
    source.add_paragraph("Art. 1º Esta norma define a tabela de valores aplicáveis.")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Categoria"
    table.cell(0, 1).text = "Valor oficial"
    buffer = io.BytesIO()
    source.save(buffer)
    parsed = _extract_docx(buffer.getvalue()).decode("utf-8")
    assert "Art. 1º" in parsed
    assert "Categoria | Valor oficial" in parsed
